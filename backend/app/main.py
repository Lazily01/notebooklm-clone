import os
os.environ["PROTOCOL_BUFFERS_PYTHON_IMPLEMENTATION"] = "python"

from fastapi import FastAPI, UploadFile, File, HTTPException, Body, Request, Depends
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import Response, FileResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel
from typing import List, Optional, Any, Dict
from pandas import DataFrame
from datetime import datetime
import os
import base64
import uuid
import json
import time
import hashlib
from dotenv import load_dotenv
from pathlib import Path

from app.services.document_parser import DocumentParser
from app.services.llm_service import LLMService
from app.services.audio_service import AudioService
from app.services.tts_service import tts_service
from app.services.rag_service import rag_service
from app.services.token_tracker import token_tracker
from app.services.web_search_service import web_search_service
from app.services.web_scraper_service import web_scraper_service
from app.database import init_db, get_db, Document, Conversation, Message, SessionLocal
from sqlalchemy.orm import Session

load_dotenv()

app = FastAPI(title="NotebookLM Clone", version="1.0.0")

# 创建静态文件目录
STATIC_DIR = Path("static")
STATIC_DIR.mkdir(exist_ok=True)
(STATIC_DIR / "mindmaps").mkdir(exist_ok=True)

# 挂载静态文件服务
app.mount("/static", StaticFiles(directory="static"), name="static")

# 初始化数据库
@app.on_event("startup")
async def startup_event():
    init_db()
    
    # 从数据库加载文档到内存缓存
    db = SessionLocal()
    try:
        docs = db.query(Document).all()
        for doc in docs:
            documents_store[doc.id] = {
                "id": doc.id,
                "name": doc.name,
                "type": doc.type,
                "content": doc.content,
                "summary": doc.summary,
                "chapters": [],
                "metadata": {},
                "raw_content": doc.raw_content
            }
        print(f"[Database] 从数据库加载了 {len(docs)} 个文档到内存")
    finally:
        db.close()

# 使用千问 Long 直接读取文件（Excel, Word 等）
async def parse_file_with_qwen(content: bytes, filename: str, file_type: str = "Excel") -> str:
    """使用千问 Long 读取文件内容"""
    import tempfile
    import os
    from openai import OpenAI
    
    # 将文件保存为临时文件
    with tempfile.NamedTemporaryFile(delete=False, suffix=os.path.splitext(filename)[1]) as tmp_file:
        tmp_file.write(content)
        tmp_path = tmp_file.name
    
    try:
        # 使用 OpenAI SDK 调用千问
        api_key = os.getenv("DASHSCOPE_API_KEY")
        client = OpenAI(
            api_key=api_key,
            base_url="https://dashscope.aliyuncs.com/compatible-mode/v1"
        )
        
        # 上传文件
        print(f"[DEBUG] 开始上传文件: {filename}")
        from pathlib import Path
        file_object = client.files.create(
            file=Path(tmp_path),
            purpose="file-extract"
        )
        file_id = file_object.id
        print(f"[DEBUG] 文件上传成功，file_id: {file_id}")
        
        # 使用 qwen-long 读取文件内容
        print(f"[DEBUG] 开始调用 qwen-long 读取文件")
        completion = client.chat.completions.create(
            model="qwen-long",
            messages=[
                {'role': 'system', 'content': f'你是一个 {file_type} 文档解析助手。'},
                {'role': 'system', 'content': f'fileid://{file_id}'},
                {'role': 'user', 'content': f'请读取这个 {file_type} 文件的内容，并将所有数据转换为文本格式输出。'}
            ]
        )
        
        text = completion.choices[0].message.content
        print(f"[DEBUG] 文件读取成功，内容长度: {len(text)}")
        
        return text
                
    except Exception as e:
        print(f"[ERROR] parse_file_with_qwen 总体失败: {e}")
        import traceback
        traceback.print_exc()
        raise
    finally:
        # 删除临时文件
        if os.path.exists(tmp_path):
            os.remove(tmp_path)

# 为了兼容，保留 Excel 的别名
async def parse_excel_with_qwen(content: bytes, filename: str) -> str:
    """使用千问 Long 读取 Excel 文件内容"""
    return await parse_file_with_qwen(content, filename, "Excel")

# CORS 配置
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# 服务实例
document_parser = DocumentParser()
llm_service = LLMService()
audio_service = AudioService()
# rag_service 已从 app.services.rag_service 导入

# 存储已上传的文档内容
documents_store: dict = {}

# Token 使用统计
token_stats = {
    # LLM 对话统计
    "total_input_tokens": 0,
    "total_output_tokens": 0,
    "total_tokens": 0,
    "chat_count": 0,
    
    # 文档上传统计
    "file_upload_count": 0,
    "summary_tokens": 0,  # 摘要生成消耗的 tokens
    "embedding_tokens": 0,  # 向量化消耗的 tokens
    "total_documents": 0,  # 总文档数
    "total_chunks": 0,  # 总片段数
    
    # RAG 检索统计
    "rag_query_count": 0,  # RAG 查询次数
    "rag_embedding_tokens": 0,  # 查询向量化消耗的 tokens
    "dashvector_queries": 0,  # DashVector 查询次数
    
    # 成本统计
    "saved_tokens": 0,  # 节省的 tokens
    "total_cost": 0.0,  # 总成本（元）
    "saved_cost": 0.0,  # 节省的成本（元）
    
    # 最近一次聊天
    "last_chat": {
        "input_tokens": 0,
        "output_tokens": 0,
        "total_tokens": 0,
        "saved_tokens": 0,
        "cost": 0.0,
        "saved_cost": 0.0
    }
}


class ChatRequest(BaseModel):
    message: str
    document_ids: List[str] = []


class ChatResponse(BaseModel):
    response: str
    sources: List[dict] = []
    suggestions: List[str] = []  # 联想问题


class DocumentInfo(BaseModel):
    id: str
    name: str
    type: str
    content_preview: str
    url: str = None  # 网页文档的 URL
    parsing_status: str = "ready"  # qwen-long 解析状态: processing, ready, error


class VideoOverviewResponse(BaseModel):
    doc_id: str
    title: str
    outline: List[Any]


class VideoOverviewRequest(BaseModel):
    length: str = "short"  # 'short' | 'default'
    style: str = "documentary"  # 'cartoon' | 'documentary' | 'narrative' | 'short'


@app.get("/")
async def root():
    return {"message": "NotebookLM Clone API", "status": "running"}


@app.post("/video/overview/{doc_id}", response_model=VideoOverviewResponse)
async def generate_video_overview(doc_id: str, request: VideoOverviewRequest = None):
    """基于文档生成视频概述大纲 + 豆包配图"""
    if doc_id not in documents_store:
        raise HTTPException(status_code=404, detail="文档不存在")

    doc = documents_store[doc_id]
    content = doc.get("content") or ""
    if not content:
        raise HTTPException(status_code=400, detail="文档内容为空，无法生成视频概述")

    # 解析参数
    length = request.length if request else "short"
    style = request.style if request else "documentary"
    section_count = 5 if length == "short" else 8

    try:
        # 1. 生成大纲
        outline = await llm_service.generate_video_outline(content, section_count, style)
        
        # 2. 用豆包为每个章节生成配图（测试模式：只生成前2张）
        import asyncio
        
        MAX_IMAGES = 2  # 测试模式：只生成2张图片
        
        for i, section in enumerate(outline):
            prompt = section.get("image_prompt", "")
            if prompt and i < MAX_IMAGES:  # 测试模式：只生成前2张
                print(f"[Video] 生成第 {i+1} 章配图...")
                try:
                    # 使用豆包 Seedream 生成配图（16:9 横屏）
                    url = await llm_service._generate_full_infographic(prompt, size="2560x1440")
                    if url:
                        section["image_url"] = url
                        print(f"[Video] 第 {i+1} 章配图成功")
                    else:
                        section["image_url"] = None
                except Exception as img_err:
                    print(f"[Video] 第 {i+1} 章配图失败: {img_err}")
                    section["image_url"] = None
                # 避免速率限制
                await asyncio.sleep(1)
        
        # 缓存大纲用于后续视频合成
        documents_store[doc_id]["video_outline"] = outline
        
        return VideoOverviewResponse(
            doc_id=doc_id,
            title=doc.get("name", "文档"),
            outline=outline
        )
    except Exception as e:
        print(f"[Video] 生成视频概述失败: {e}")
        raise HTTPException(status_code=500, detail=str(e))


# 视频存储（重启后端时清空）
video_store = {}  # {video_id: {"file_path": str, "title": str, "doc_id": str}}

@app.post("/video/generate/{doc_id}")
async def generate_video(doc_id: str, voice: str = "Cherry"):
    """根据已生成的大纲合成视频（需要先调用 /video/overview）"""
    from app.services.video_service import VideoService
    import uuid
    from pathlib import Path
    
    if doc_id not in documents_store:
        raise HTTPException(status_code=404, detail="文档不存在")
    
    # 检查是否有缓存的大纲
    doc = documents_store[doc_id]
    outline = doc.get("video_outline")
    
    if not outline:
        raise HTTPException(status_code=400, detail="请先生成视频概述大纲")
    
    try:
        # 创建视频存储目录
        video_dir = Path("videos")
        video_dir.mkdir(exist_ok=True)
        
        # 生成视频
        video_service = VideoService(llm_service)
        temp_video_path = await video_service.generate_video_from_outline(outline, doc_id, voice)
        
        if temp_video_path and os.path.exists(temp_video_path):
            # 生成唯一视频ID
            video_id = str(uuid.uuid4())[:8]
            
            # 移动到永久存储目录
            video_store_path = video_dir / f"{video_id}.mp4"
            import shutil
            shutil.move(temp_video_path, video_store_path)
            
            # 保存到视频存储
            video_store[video_id] = {
                "file_path": str(video_store_path),
                "title": doc.get("name", "视频概述"),
                "doc_id": doc_id
            }
            
            # 返回视频信息
            return {
                "video_id": video_id,
                "title": doc.get("name", "视频概述")
            }
        else:
            raise HTTPException(status_code=500, detail="视频生成失败")
    except Exception as e:
        print(f"[Video] 视频合成失败: {e}")
        raise HTTPException(status_code=500, detail=str(e))

@app.get("/video/{video_id}")
async def get_video(video_id: str, request: Request):
    """获取视频文件（支持 Range 请求以实现进度条拖动）"""
    from starlette.responses import StreamingResponse
    
    if video_id not in video_store:
        raise HTTPException(status_code=404, detail="视频不存在")
    
    video_info = video_store[video_id]
    file_path = video_info["file_path"]
    
    if not os.path.exists(file_path):
        raise HTTPException(status_code=404, detail="视频文件不存在")
    
    file_size = os.path.getsize(file_path)
    
    # 检查是否有 Range 请求头
    range_header = request.headers.get("range") if request else None
    
    if range_header:
        # 解析 Range 头
        range_match = range_header.replace("bytes=", "").split("-")
        start = int(range_match[0]) if range_match[0] else 0
        end = int(range_match[1]) if range_match[1] else file_size - 1
        
        # 确保范围有效
        if start >= file_size:
            raise HTTPException(status_code=416, detail="Range Not Satisfiable")
        
        end = min(end, file_size - 1)
        content_length = end - start + 1
        
        def iter_file():
            with open(file_path, "rb") as f:
                f.seek(start)
                remaining = content_length
                while remaining > 0:
                    chunk_size = min(8192, remaining)
                    data = f.read(chunk_size)
                    if not data:
                        break
                    remaining -= len(data)
                    yield data
        
        return StreamingResponse(
            iter_file(),
            status_code=206,
            media_type="video/mp4",
            headers={
                "Content-Range": f"bytes {start}-{end}/{file_size}",
                "Accept-Ranges": "bytes",
                "Content-Length": str(content_length),
            }
        )
    else:
        # 没有 Range 请求，返回完整文件
        return FileResponse(
            file_path,
            media_type="video/mp4",
            filename=f"{video_info['title']}.mp4",
            headers={"Accept-Ranges": "bytes"}
        )


@app.post("/upload", response_model=DocumentInfo)
async def upload_document(file: UploadFile = File(...), db: Session = Depends(get_db)):
    """上传并解析文档（PDF、Word、音频、图片）"""
    try:
        # 检查文档数量限制
        MAX_DOCUMENTS = 50  # 最大文档数量
        current_count = db.query(Document).count()
        if current_count >= MAX_DOCUMENTS:
            raise HTTPException(
                status_code=400, 
                detail=f"文档数量已达上限（{MAX_DOCUMENTS}个），请删除部分文档后再上传"
            )
        # 读取文件内容
        content = await file.read()
        filename = file.filename
        
        # 计算文件哈希，检查是否已存在
        import hashlib
        file_hash = hashlib.md5(content).hexdigest()
        
        # 查询数据库是否已有相同文件
        existing_doc = db.query(Document).filter(Document.file_hash == file_hash).first()
        if existing_doc:
            print(f"[Cache] 文件已存在，使用缓存: {existing_doc.id}")
            # 同步到内存缓存
            if existing_doc.id not in documents_store:
                documents_store[existing_doc.id] = {
                    "id": existing_doc.id,
                    "name": existing_doc.name,
                    "type": existing_doc.type,
                    "content": existing_doc.content,
                    "summary": existing_doc.summary,
                    "chapters": [],
                    "metadata": {},
                    "raw_content": existing_doc.raw_content
                }
            return DocumentInfo(
                id=existing_doc.id,
                name=existing_doc.name,
                type=existing_doc.type,
                content_preview=existing_doc.content_preview or "文档已上传"
            )
        
        # 根据文件类型解析
        is_ocr_pdf = False  # 标记是否为 OCR 处理的 PDF
        need_page_estimation = False  # 标记是否需要页码估算
        
        if filename.endswith('.pdf'):
            # PDF：先尝试提取文本
            result = document_parser.parse_pdf(content)
            text = result["text"]
            
            # 如果有需要 OCR 的页面（乱码或扫描件）
            if result.get("images") and len(result["images"]) > 0:
                print(f"[PDF] 检测到 {len(result['images'])} 页需要 OCR，使用 Qwen-Long 解析整个 PDF")
                
                # 上传整个 PDF 到 Qwen-Long
                file_id = await llm_service.upload_file_and_get_id(content, filename)
                print(f"[PDF] 上传成功，file-id: {file_id}")
                
                # 使用 Qwen-Long 提取文本
                extracted_text = await llm_service.extract_text_from_file(file_id)
                print(f"[PDF] Qwen-Long 解析完成，文本长度: {len(extracted_text)}")
                
                # 使用 Qwen-Long 提取的文本
                text = extracted_text
                is_ocr_pdf = True  # 标记为 OCR PDF
                need_page_estimation = True  # 需要估算页码
            
            doc_type = "pdf"
        elif filename.lower().endswith(('.png', '.jpg', '.jpeg', '.gif', '.bmp', '.webp')):
            # 图片文件：使用 Qwen-Long（支持图片，比 VL-Max 便宜 16 倍）
            img_base64 = base64.b64encode(content).decode('utf-8')
            
            # 检测图片类型
            import imghdr
            img_type = imghdr.what(None, content) or 'jpeg'
            
            # 使用 Qwen-Long 识别图片
            print(f"[Image] 使用 Qwen-Long 识别图片，类型: {img_type}")
            text = await llm_service.recognize_image(img_base64, img_type)
            doc_type = "image"
        elif filename.endswith(('.docx', '.doc')):
            # Word 文件：本地解析优先
            result = document_parser.parse_docx(content)
            if result["quality"] in ["good", "medium"]:
                text = result["text"]
                print(f"[Word] 本地解析成功，质量: {result['quality']}")
            else:
                # 质量差，使用 Qwen-Long
                print(f"[Word] 本地解析质量差，使用 Qwen-Long")
                text = await parse_file_with_qwen(content, filename, "Word")
            doc_type = "word"
            # Word 不需要页码估算（不显示引用）
        elif filename.endswith(('.xlsx', '.xls')):
            # Excel 文件：本地解析优先
            result = document_parser.parse_excel(content)
            if result["quality"] in ["good", "medium"]:
                text = result["text"]
                print(f"[Excel] 本地解析成功，质量: {result['quality']}, {result.get('total_rows', 0)} 行")
            else:
                # 质量差或有复杂公式，使用 Qwen-Long
                print(f"[Excel] 本地解析质量差，使用 Qwen-Long")
                text = await parse_excel_with_qwen(content, filename)
            doc_type = "excel"
            # Excel 不需要页码估算（不显示引用）
        elif filename.endswith('.md'):
            # Markdown：增强解析（提取结构）
            result = document_parser.parse_markdown(content)
            text = result["text"]
            print(f"[MD] 解析成功，标题: {result['structure']['headers']} 个，代码块: {result['structure']['code_blocks']} 个")
            doc_type = "markdown"
        elif filename.endswith(('.epub', '.mobi')):
            # 电子书：尝试使用 ebooklib 提取正文
            try:
                from ebooklib import epub
            except ImportError:
                raise HTTPException(status_code=400, detail="解析 EPUB/MOBI 需要安装 ebooklib，请先安装依赖")

            from bs4 import BeautifulSoup
            from io import BytesIO

            try:
                book = epub.read_epub(BytesIO(content))
            except Exception as e:
                raise HTTPException(status_code=400, detail=f"电子书解析失败: {e}")

            items = []
            
            # 获取所有项目
            all_items = list(book.get_items())
            
            for idx, item in enumerate(all_items):
                item_type = item.get_type()
                
                # 尝试多种内容类型
                if item_type == epub.EpubHtml or item_type == 9:  # 9 是 ITEM_DOCUMENT
                    try:
                        html = item.get_content().decode('utf-8', errors='ignore')
                        soup = BeautifulSoup(html, 'html.parser')
                        
                        # 移除 script 和 style 标签
                        for script in soup(["script", "style"]):
                            script.decompose()
                        
                        text_part = soup.get_text(separator='\n', strip=True)
                        
                        # 清理空行
                        lines = [line.strip() for line in text_part.split('\n') if line.strip()]
                        text_part = '\n'.join(lines)
                        
                        if text_part and len(text_part) > 10:  # 至少10个字符
                            items.append(text_part)
                    except Exception as e:
                        pass  # 静默处理失败的项目
            
            text = "\n\n".join(items)
            print(f"[EPUB] 解析完成，文本长度: {len(text)}")
            
            if not text or len(text) < 50:
                raise HTTPException(status_code=400, detail="EPUB 文件解析失败：未能提取到有效文本内容")
            
            doc_type = "ebook"
        elif filename.endswith(('.mp3', '.wav', '.m4a', '.ogg', '.flac')):
            # 先生成 doc_id，用于异步任务
            doc_id = f"doc_{uuid.uuid4().hex[:12]}"
            # 启动异步转录
            text = await audio_service.start_transcription(content, filename, doc_id)
            doc_type = "audio"
            
            # 存储文档（转录中）
            documents_store[doc_id] = {
                "id": doc_id,
                "name": filename,
                "type": doc_type,
                "content": text,
                "raw_content": base64.b64encode(content).decode('utf-8'),
                "transcription_status": "processing"
            }
            
            return DocumentInfo(
                id=doc_id,
                name=filename,
                type=doc_type,
                content_preview="正在转录音频..."
            )
        elif filename.endswith('.txt'):
            # TXT：智能编码检测
            result = document_parser.parse_txt(content)
            text = result["text"]
            print(f"[TXT] 解析成功，编码: {result['encoding']}, 置信度: {result.get('confidence', 0):.2f}")
            doc_type = "text"
        else:
            raise HTTPException(status_code=400, detail="不支持的文件格式")
        
        print(f"[DEBUG] 文档内容长度: {len(text) if text else 0}")
        
        # RAG 方案不需要上传 file_id，直接使用本地解析的文本
        # 节省上传时间和存储空间
        
        # 提取章节结构（EPUB）
        chapters = []
        if doc_type == "epub":
            chapters = await extract_epub_chapters(content)
        
        # 不在上传时生成摘要，改为按需生成（节省成本和时间）
        summary = None  # 摘要将在用户点击"摘要"按钮时按需生成
        
        # 提取文档元数据（0成本，快速）
        import re
        metadata = {
            "word_count": len(text),
            "char_count": len(text),
            "preview": text[:500] if text else "",  # 前500字预览
            # 提取标题结构（Markdown格式）
            "headings": re.findall(r'^#{1,6}\s+(.+)$', text, re.MULTILINE) if text else [],
            # 快速关键词（简单实现，不调用LLM）
            "has_tables": "表" in text or "|" in text,
            "has_code": "```" in text or "def " in text or "function " in text,
        }
        
        # 生成文档ID并存储（使用UUID避免冲突）
        doc_id = f"doc_{uuid.uuid4().hex[:12]}"
        
        # 添加到向量库（RAG）- 所有文档都向量化
        rag_stats = {"chunks_count": 0, "tokens_used": 0}
        
        print(f"[RAG] 开始向量化文档 ({len(text)} 字符)")
        try:
            if doc_type == "pdf" and not is_ocr_pdf:
                # 正常 PDF 文件：解析并记录真实页码（从字节流）
                chunks_with_pages = rag_service.parse_pdf_with_pages(pdf_content=content)
                if chunks_with_pages:
                    rag_stats = await rag_service.add_document_with_pages(
                        doc_id=doc_id,
                        chunks_with_pages=chunks_with_pages,
                        metadata={
                            "filename": filename,
                            "doc_type": doc_type
                        }
                    )
                    print(f"[RAG] PDF 文档 {doc_id} 已添加到向量库（带真实页码）")
                else:
                    print(f"[RAG] PDF 解析失败，跳过向量化")
            else:
                # OCR PDF 或非 PDF 文件：使用普通文本切分
                rag_stats = await rag_service.add_document(
                    doc_id=doc_id,
                    text=text,
                    metadata={
                        "filename": filename,
                        "doc_type": doc_type
                    },
                    estimate_pages=need_page_estimation  # 需要估算页码的文档
                )
                if need_page_estimation:
                    print(f"[RAG] {doc_type.upper()} 文档 {doc_id} 已添加到向量库（普通文本切分 + 估算页码）")
                else:
                    print(f"[RAG] 文档 {doc_id} 已添加到向量库")
            
            # 更新统计
            token_stats["embedding_tokens"] += rag_stats["tokens_used"]
            token_stats["total_chunks"] += rag_stats["chunks_count"]
            token_stats["total_documents"] += 1
        except Exception as e:
            print(f"[RAG] 向量化失败: {e}")
            # 不阻塞上传流程
        
        # 保存到数据库（替代 documents_store）
        content_preview = text[:200] + "..." if text and len(text) > 200 else (text or "文档已上传")
        
        db_document = Document(
            id=doc_id,
            name=filename,
            type=doc_type,
            file_hash=file_hash,
            content=text,
            raw_content=base64.b64encode(content).decode('utf-8'),
            summary=summary,
            content_preview=content_preview,
            parsing_status="ready"
        )
        db.add(db_document)
        db.commit()
        db.refresh(db_document)
        
        print(f"[Database] 文档已保存: {doc_id}")
        
        # 同时保存到内存缓存（兼容现有代码）
        documents_store[doc_id] = {
            "id": doc_id,
            "name": filename,
            "type": doc_type,
            "content": text,
            "summary": summary,
            "chapters": chapters,
            "metadata": metadata,
            "raw_content": base64.b64encode(content).decode('utf-8')
        }
        
        # 统计文件上传
        token_stats["file_upload_count"] += 1
        
        return DocumentInfo(
            id=doc_id,
            name=filename,
            type=doc_type,
            content_preview=content_preview
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@app.get("/documents", response_model=List[DocumentInfo])
async def list_documents(db: Session = Depends(get_db)):
    """获取所有已上传的文档"""
    docs = db.query(Document).all()
    return [
        DocumentInfo(
            id=doc.id,
            name=doc.name,
            type=doc.type,
            content_preview=doc.content_preview or "文档已上传",
            parsing_status=doc.parsing_status or "ready"
        )
        for doc in docs
    ]


# RAG 方案不需要检查解析状态，文档上传后立即可用


@app.get("/token-stats")
async def get_token_stats():
    """查看 Token 使用统计（完整版，包含 RAG 统计）"""
    
    # 计算向量化成本（text-embedding-v4: 0.00025元/千token，Batch 价格）
    embedding_cost = (token_stats["embedding_tokens"] / 1000) * 0.00025
    rag_embedding_cost = (token_stats["rag_embedding_tokens"] / 1000) * 0.0005
    
    # 计算摘要成本
    summary_cost = (token_stats["summary_tokens"] / 1000) * 0.0005
    
    # DashVector 成本（估算）
    dashvector_storage_cost = (token_stats["total_chunks"] / 1000) * 0.0003 * 30  # 月成本
    dashvector_query_cost = (token_stats["dashvector_queries"] / 1000) * 0.002
    
    return {
        # LLM 对话统计
        "llm_usage": {
            "total_input_tokens": token_stats["total_input_tokens"],
            "total_output_tokens": token_stats["total_output_tokens"],
            "total_tokens": token_stats["total_tokens"],
            "chat_count": token_stats["chat_count"],
            "saved_tokens": token_stats["saved_tokens"],
        },
        
        # 文档上传统计
        "document_stats": {
            "total_documents": token_stats["total_documents"],
            "total_chunks": token_stats["total_chunks"],
            "summary_tokens": token_stats["summary_tokens"],
            "embedding_tokens": token_stats["embedding_tokens"],
        },
        
        # RAG 检索统计
        "rag_stats": {
            "query_count": token_stats["rag_query_count"],
            "embedding_tokens": token_stats["rag_embedding_tokens"],
            "dashvector_queries": token_stats["dashvector_queries"],
        },
        
        # 成本明细（元）
        "cost_breakdown": {
            "llm_chat": round(token_stats["total_cost"], 6),
            "summary_generation": round(summary_cost, 6),
            "document_embedding": round(embedding_cost, 6),
            "query_embedding": round(rag_embedding_cost, 6),
            "dashvector_storage": round(dashvector_storage_cost, 6),
            "dashvector_query": round(dashvector_query_cost, 6),
            "total": round(token_stats["total_cost"] + summary_cost + embedding_cost + rag_embedding_cost + dashvector_storage_cost + dashvector_query_cost, 6),
            "saved": round(token_stats["saved_cost"], 6),
        },
        
        # 效率指标
        "efficiency": {
            "avg_tokens_per_chat": round(token_stats["total_tokens"] / token_stats["chat_count"], 2) if token_stats["chat_count"] > 0 else 0,
            "avg_chunks_per_doc": round(token_stats["total_chunks"] / token_stats["total_documents"], 2) if token_stats["total_documents"] > 0 else 0,
            "savings_rate": f"{round(token_stats['saved_tokens'] / (token_stats['saved_tokens'] + token_stats['total_input_tokens']) * 100, 2)}%" if (token_stats['saved_tokens'] + token_stats['total_input_tokens']) > 0 else "0%"
        },
        
        # 最近一次聊天
        "last_chat": token_stats["last_chat"]
    }


@app.get("/qwen-files")
async def list_qwen_files():
    """查看阿里云上传的所有文件"""
    try:
        from openai import OpenAI
        client = OpenAI(
            api_key=os.getenv("DASHSCOPE_API_KEY"),
            base_url="https://dashscope.aliyuncs.com/compatible-mode/v1"
        )
        files = client.files.list()
        
        file_list = []
        for file in files.data:
            file_list.append({
                "id": file.id,
                "filename": file.filename,
                "bytes": file.bytes,
                "created_at": file.created_at,
                "status": file.status,
                "purpose": file.purpose
            })
        
        return {
            "total": len(file_list),
            "files": file_list
        }
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"获取文件列表失败: {str(e)}")


@app.delete("/qwen-files/cleanup")
async def cleanup_qwen_files():
    """清理阿里云上的所有文件"""
    try:
        from openai import OpenAI
        client = OpenAI(
            api_key=os.getenv("DASHSCOPE_API_KEY"),
            base_url="https://dashscope.aliyuncs.com/compatible-mode/v1"
        )
        files = client.files.list()
        
        deleted_count = 0
        failed_count = 0
        
        for file in files.data:
            try:
                client.files.delete(file.id)
                print(f"[Cleanup] 已删除文件: {file.filename} ({file.id})")
                deleted_count += 1
            except Exception as e:
                print(f"[Cleanup] 删除失败: {file.filename} - {e}")
                failed_count += 1
        
        return {
            "message": "清理完成",
            "deleted": deleted_count,
            "failed": failed_count
        }
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"清理失败: {str(e)}")


@app.delete("/documents/{doc_id}")
async def delete_document(doc_id: str, db: Session = Depends(get_db)):
    """删除文档"""
    # 先检查数据库（而不是内存）
    db_doc = db.query(Document).filter(Document.id == doc_id).first()
    if not db_doc:
        raise HTTPException(status_code=404, detail="文档不存在")
    
    # 1. 删除向量库中的文档
    try:
        await rag_service.delete_document(doc_id)
        print(f"[RAG] 已从向量库删除文档: {doc_id}")
    except Exception as e:
        print(f"[RAG] 删除向量失败: {e}")
    
    # 2. 删除数据库记录
    try:
        db.delete(db_doc)
        db.commit()
        print(f"[Database] 已从数据库删除文档: {doc_id}")
    except Exception as e:
        print(f"[Database] 删除数据库记录失败: {e}")
        raise HTTPException(status_code=500, detail=f"删除失败: {str(e)}")
    
    # 3. 删除内存缓存（如果存在）
    if doc_id in documents_store:
        del documents_store[doc_id]
        print(f"[Memory] 已从内存删除文档: {doc_id}")
    
    return {"message": "文档已删除"}


@app.get("/token-stats")
async def get_token_stats():
    """获取 Token 使用统计（透明化成本）"""
    summary = token_tracker.get_summary()
    return summary


@app.post("/token-stats/reset")
async def reset_token_stats():
    """重置 Token 统计"""
    token_tracker.reset()
    return {"message": "统计已重置"}


# ==================== 联网搜索 ====================

class WebSearchRequest(BaseModel):
    """联网搜索请求"""
    query: str
    max_pages: int = 3  # 最多爬取几个网页


@app.post("/web-search-preview")
async def web_search_preview(request: dict):
    """
    联网搜索预览（只搜索，不爬取，不导入）
    返回搜索结果列表供用户查看
    """
    query = request.get("query", "")
    max_results = request.get("max_results", 5)
    
    if not query:
        raise HTTPException(status_code=400, detail="搜索关键词不能为空")
    
    print(f"[WebSearch] 搜索预览: {query}")
    
    try:
        # 只搜索，返回结果列表
        results = await web_search_service.search_web(query, count=max_results)
        
        if not results:
            raise HTTPException(status_code=404, detail="未找到搜索结果")
        
        print(f"[WebSearch] 找到 {len(results)} 个结果")
        
        return {
            "query": query,
            "results": results,
            "count": len(results)
        }
        
    except Exception as e:
        print(f"[WebSearch] 搜索失败: {e}")
        raise HTTPException(status_code=500, detail=f"搜索失败: {str(e)}")


@app.post("/web-search-import")
async def web_search_import(request: dict, db: Session = Depends(get_db)):
    """
    导入网页（爬取并添加到文档库）
    """
    urls = request.get("urls", [])
    
    if not urls:
        raise HTTPException(status_code=400, detail="URL 列表不能为空")
    
    print(f"[WebSearch] 开始导入 {len(urls)} 个网页")
    
    added_docs = []
    
    try:
        for i, url in enumerate(urls, 1):
            print(f"[WebSearch] 爬取 {i}/{len(urls)}: {url}")
            
            # 爬取网页内容
            webpage = await web_search_service.fetch_webpage_content(url)
            
            if not webpage:
                print(f"[WebSearch] 跳过爬取失败的网页: {url}")
                continue
            
            # 生成文档 ID 和哈希
            doc_id = f"doc_{int(time.time() * 1000)}_{i}"
            file_hash = hashlib.md5(url.encode()).hexdigest()
            
            # 检查是否已存在
            existing_doc = db.query(Document).filter(Document.file_hash == file_hash).first()
            if existing_doc:
                print(f"[WebSearch] 网页已存在，跳过: {url}")
                continue
            
            # 添加到向量库
            metadata = {
                "doc_id": doc_id,
                "type": "webpage",
                "url": url
            }
            
            await rag_service.add_document(
                doc_id=doc_id,
                text=webpage["content"],
                metadata=metadata,
                estimate_pages=False
            )
            
            # 保存到数据库
            db_doc = Document(
                id=doc_id,
                name=webpage["title"][:100],
                type="webpage",
                content=webpage["content"][:500],
                raw_content="",
                file_hash=file_hash,
                parsing_status="ready"
            )
            db.add(db_doc)
            db.commit()
            
            # 添加到内存
            documents_store[doc_id] = {
                "id": doc_id,
                "name": webpage["title"],
                "type": "webpage",
                "content": webpage["content"][:500],
                "url": url,
                "raw_content": webpage["content"]  # 保存完整内容用于预览
            }
            
            added_docs.append({
                "id": doc_id,
                "name": webpage["title"],
                "url": url
            })
            
            print(f"[WebSearch] 网页 {doc_id} 已添加")
        
        return {
            "message": f"成功导入 {len(added_docs)} 个网页",
            "documents": added_docs
        }
        
    except Exception as e:
        print(f"[WebSearch] 导入失败: {e}")
        raise HTTPException(status_code=500, detail=f"导入失败: {str(e)}")


@app.get("/documents/{doc_id}/content")
async def get_document_content(doc_id: str):
    """获取文档完整内容（用于音频转录文本预览和网页预览）"""
    if doc_id not in documents_store:
        raise HTTPException(status_code=404, detail="文档不存在")
    
    doc = documents_store[doc_id]
    return {
        "content": doc.get("raw_content") or doc.get("content", ""),
        "name": doc.get("name", ""),
        "url": doc.get("url", ""),
        "type": doc.get("type", "")
    }


@app.get("/documents/{doc_id}/file")
async def get_document_file(doc_id: str):
    """获取文档原始文件"""
    if doc_id not in documents_store:
        raise HTTPException(status_code=404, detail="文档不存在")
    
    doc = documents_store[doc_id]
    raw_content = base64.b64decode(doc["raw_content"])
    
    # 根据文件类型设置 Content-Type
    content_type = "application/pdf" if doc["type"] == "pdf" else "application/octet-stream"
    
    # 对文件名进行 URL 编码以支持中文
    from urllib.parse import quote
    encoded_filename = quote(doc['name'])
    
    return Response(
        content=raw_content,
        media_type=content_type,
        headers={
            "Content-Disposition": f"inline; filename*=UTF-8''{encoded_filename}",
            "Cache-Control": "public, max-age=3600",  # 缓存1小时
        }
    )


@app.post("/chat")
async def chat(request: ChatRequest):
    """非流式聊天接口（用于闪卡详细解释等场景）"""
    print(f"[Chat] 收到聊天请求: {request.message}")
    
    try:
        # 获取文档列表
        doc_ids = request.document_ids if request.document_ids else list(documents_store.keys())
        print(f"[Chat] 文档列表: {doc_ids}")
        
        if not doc_ids:
            raise HTTPException(status_code=400, detail="请先上传文档")
        
        # 使用 RAG 检索相关片段
        print(f"[RAG] 使用 RAG 检索")
        rag_result = {"chunks": [], "tokens_used": 0}
        try:
            result = await rag_service.search(
                query=request.message,
                top_k=8,
                filter_doc_ids=doc_ids
            )
            if result:
                rag_result = result
        except Exception as e:
            print(f"[RAG] 检索失败: {e}")
        
        # 构建上下文
        relevant_chunks = rag_result.get("chunks", [])
        print(f"[Chat] RAG 检索到 {len(relevant_chunks)} 个片段")
        
        context_parts = []
        sources = []
        seen_pages = set()
        
        for idx, chunk in enumerate(relevant_chunks, 1):
            doc_id = chunk.get("doc_id", "")
            doc = documents_store.get(doc_id)
            doc_name = doc.get("name", "未知文档") if doc else "未知文档"
            chunk_text = chunk.get("text", "")
            page = chunk.get("page", "未知")
            
            page_key = f"{doc_id}_{page}"
            if page_key not in seen_pages:
                seen_pages.add(page_key)
                # 使用和流式对话一样的格式
                context_parts.append(f"===== 第 {page} 页 =====\n文档: {doc_name}\n内容: {chunk_text}\n⚠️ 引用此内容时必须使用: [{page}]")
                sources.append({
                    "doc_id": doc_id,
                    "doc_name": doc_name,
                    "page": page,
                    "text": chunk_text[:200]
                })
        
        context = "\n\n".join(context_parts) if context_parts else "无相关内容"
        
        # 调用 LLM
        result = await llm_service.chat(
            question=request.message,
            context=context
        )
        
        return ChatResponse(
            response=result["response"],
            sources=sources,
            suggestions=result.get("suggestions", [])
        )
    
    except Exception as e:
        print(f"[Chat] 对话失败: {e}")
        import traceback
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/chat/stream")
async def chat_stream(request: ChatRequest):
    """流式聊天接口（SSE）"""
    print(f"[Stream] 收到流式聊天请求: {request.message}")
    
    async def generate():
        try:
            print(f"[Stream] 开始生成响应...")
            # 获取文档列表
            doc_ids = request.document_ids if request.document_ids else list(documents_store.keys())
            print(f"[Stream] 文档列表: {doc_ids}")
            
            if not doc_ids:
                yield f"data: {json.dumps({'error': '请先上传文档'})}\n\n"
                return
            
            # 使用 RAG 检索相关片段（所有文档都使用 RAG）
            print(f"[RAG] 使用 RAG 检索")
            rag_result = {"chunks": [], "tokens_used": 0}
            try:
                result = await rag_service.search(
                    query=request.message,
                    top_k=8,  # 增加到 8 个，提高覆盖率
                    filter_doc_ids=doc_ids
                )
                if result:
                    rag_result = result
            except Exception as e:
                print(f"[RAG] 检索失败: {e}")
                import traceback
                traceback.print_exc()
            
            # 构建上下文
            relevant_chunks = rag_result.get("chunks", [])
            print(f"[Stream] RAG 检索到 {len(relevant_chunks)} 个片段")
            
            if relevant_chunks:
                context_parts = []
                seen_pages = set()  # 记录已添加的页码，避免重复
                
                for idx, chunk in enumerate(relevant_chunks, 1):
                    doc_id = chunk.get("doc_id", "")
                    doc = documents_store.get(doc_id)
                    doc_name = doc.get("name", "未知文档") if doc else "未知文档"
                    doc_type = doc.get("type", "") if doc else ""
                    chunk_text = chunk.get("text", "")
                    page = chunk.get("page", "未知")
                    
                    # 去重：如果这一页已经添加过，跳过
                    page_key = f"{doc_id}_{page}"
                    if page_key in seen_pages:
                        continue
                    seen_pages.add(page_key)
                    
                    # 只有 PDF 文件才添加引用提示
                    if doc_type == "pdf":
                        # 使用页码作为引用编号
                        context_parts.append(f"===== 第 {page} 页 =====\n文档: {doc_name}\n内容: {chunk_text}\n⚠️ 引用此内容时必须使用: [{page}]")
                    else:
                        # Word/Excel 等不添加引用
                        context_parts.append(f"===== 来源 =====\n文档: {doc_name}\n内容: {chunk_text}")
                
                context = "\n\n".join(context_parts)
                print(f"[Stream] 使用 RAG 上下文，去重后 {len(context_parts)} 个页面，长度: {len(context)}")
            else:
                # 降级：使用摘要
                print(f"[Stream] RAG 未找到相关内容，降级使用摘要")
                context_parts = []
                for doc_id in doc_ids:
                    if doc_id in documents_store:
                        doc = documents_store[doc_id]
                        summary = doc.get("summary") or ""
                        content = doc.get("content") or ""
                        print(f"[Stream] 文档 {doc_id}: 摘要长度={len(summary)}, 内容长度={len(content)}")
                        if summary:
                            context_parts.append(f"[文档: {doc['name']}]\n{summary}")
                        elif content:
                            # 如果没有摘要，使用内容的前 1000 字符
                            context_parts.append(f"[文档: {doc['name']}]\n{content[:1000]}")
                context = "\n\n".join(context_parts)
                print(f"[Stream] 使用摘要/内容上下文，长度: {len(context)}")
            
            # 计算并记录向量检索内容的 token 消耗
            context_tokens = len(context) // 2  # 粗略估算：中文约2字符/token
            context_cost = context_tokens * 0.0005 / 1000  # qwen-long 输入价格
            
            # 计算全文档的 token 数（如果不用 RAG）
            total_doc_chars = sum(len(documents_store[doc_id].get("content", "")) 
                                 for doc_id in doc_ids 
                                 if doc_id in documents_store)
            total_doc_tokens = total_doc_chars // 2
            total_doc_cost = total_doc_tokens * 0.0005 / 1000
            
            # 计算节省的 tokens
            saved_tokens = total_doc_tokens - context_tokens
            saved_cost = total_doc_cost - context_cost
            saved_percentage = (saved_tokens / total_doc_tokens * 100) if total_doc_tokens > 0 else 0
            
            # Token 统计
            print(f"[Token] ✅ 本次使用RAG: {context_tokens} tokens (约 {len(context)} 字符), 成本: {context_cost:.6f}元")
            if saved_tokens > 0:
                print(f"[Token] 节省: {saved_tokens} tokens ({saved_percentage:.1f}%)")
            
            # 构建消息（使用动态 Prompt）
            from app.services.prompt_manager import prompt_manager
            
            # 获取文档类型（使用第一个文档的类型）
            doc_type = "text"
            if doc_ids and doc_ids[0] in documents_store:
                doc_type = documents_store[doc_ids[0]].get("type", "text")
            
            # 使用动态 Prompt
            prompt = prompt_manager.get_chat_prompt(
                doc_type=doc_type,
                query=request.message,
                context=context
            )
            
            messages = [
                {"role": "user", "content": prompt}
            ]
            
            # 调用流式 API
            print(f"[Stream] 开始调用 LLM 流式 API...")
            stream = await llm_service._call_chat_api(messages, stream=True)
            print(f"[Stream] 流式生成器已创建")
            
            # 逐块发送（但不发送 SUGGESTIONS 部分）
            full_response = ""
            suggestions_started = False
            async for chunk in stream:
                full_response += chunk
                
                # 检查是否开始 SUGGESTIONS 部分
                if "[SUGGESTIONS]" in full_response and not suggestions_started:
                    suggestions_started = True
                    # 发送清理后的内容（替换之前的流式输出）
                    parts = full_response.split("[SUGGESTIONS]")
                    clean_content = parts[0].strip()
                    yield f"data: {json.dumps({'replace': clean_content})}\n\n"
                    # 继续读取但不发送，直到流结束
                elif not suggestions_started:
                    # 正常发送流式内容
                    yield f"data: {json.dumps({'content': chunk})}\n\n"
                # 如果已经开始 SUGGESTIONS，继续读取但不发送
            
            # 调试：检查 LLM 原始输出是否已包含引用标记
            print(f"[Stream] LLM 原始输出（前200字符）: {full_response[:200]}")
            if '[1]' in full_response or '[2]' in full_response:
                print(f"[Stream] ✅ LLM 已自动添加引用标记")
            else:
                print(f"[Stream] ❌ LLM 未添加引用标记")
            
            # 构建引用来源（详细信息，包含页码）
            sources = []
            
            # 如果有 RAG 检索结果，使用检索到的 chunks
            if relevant_chunks:
                for idx, chunk in enumerate(relevant_chunks):
                    doc_id = chunk.get("doc_id", "")
                    doc = documents_store.get(doc_id)
                    if doc:
                        # 获取真实页码（如果有）或估算
                        page = chunk.get("page")
                        chunk_id = chunk.get("chunk_id", 0)
                        
                        # 确保 page 是整数，不是字典或其他类型
                        if page is None or isinstance(page, dict):
                            # 降级：估算页码（基于 chunk_id）
                            page = (chunk_id * 2) // 3 + 1
                        elif not isinstance(page, (int, str)):
                            # 如果是其他类型，尝试转换为整数
                            try:
                                page = int(page)
                            except (ValueError, TypeError):
                                page = (chunk_id * 2) // 3 + 1
                        
                        sources.append({
                            "id": idx + 1,  # 引用编号
                            "doc_id": doc_id,
                            "name": doc.get("name", "未知文档"),
                            "type": doc.get("type", "text"),
                            "text": chunk.get("text", "")[:200],  # 预览前 200 字
                            "chunk_id": chunk.get("chunk_id", 0),
                            "page": page,  # 真实页码或估算页码（确保是 int 或 str）
                            "score": chunk.get("score", 0)  # 相似度分数
                        })
            
            # 解析联想问题
            suggestions = []
            if "[SUGGESTIONS]" in full_response:
                parts = full_response.split("[SUGGESTIONS]")
                answer = parts[0].strip()
                suggestions_text = parts[1].strip()
                suggestions = [q.strip() for q in suggestions_text.split('\n') if q.strip()]
                # 替换响应内容为不包含 SUGGESTIONS 的部分
                yield f"data: {json.dumps({'replace': answer})}\n\n"
            
            # 发送引用来源（先发送，确保前端先接收到）
            if sources:
                yield f"data: {json.dumps({'sources': sources})}\n\n"
            
            # 发送联想问题
            if suggestions:
                yield f"data: {json.dumps({'suggestions': suggestions})}\n\n"
            
            # 发送结束信号
            yield f"data: {json.dumps({'done': True})}\n\n"
        
        except Exception as e:
            import traceback
            error_detail = traceback.format_exc()
            print(f"[Stream] 错误: {e}")
            print(f"[Stream] 详细错误:\n{error_detail}")
            yield f"data: {json.dumps({'error': str(e)})}\n\n"
    
    return StreamingResponse(
        generate(),
        media_type="text/event-stream",
        headers={
            "Cache-Control": "no-cache",
            "X-Accel-Buffering": "no"
        }
    )


@app.get("/transcription/{doc_id}")
async def get_transcription_status(doc_id: str):
    """获取音频转录状态"""
    if doc_id not in documents_store:
        raise HTTPException(status_code=404, detail="文档不存在")
    
    doc = documents_store[doc_id]
    if doc["type"] != "audio":
        return {"status": "completed", "text": doc["content"]}
    
    # 获取转录状态
    status = audio_service.get_transcription_status(doc_id)
    if not status:
        return {"status": "unknown", "text": doc["content"]}
    
    # 如果转录完成，更新文档内容
    if status["status"] == "completed" and status["text"]:
        documents_store[doc_id]["content"] = status["text"]
        documents_store[doc_id]["transcription_status"] = "completed"
    
    return {
        "status": status["status"],
        "text": status["text"] if status["status"] == "completed" else doc["content"],
        "error": status.get("error")
    }


@app.post("/summarize/{doc_id}")
async def summarize_document(doc_id: str, regenerate: bool = False):
    """获取或重新生成文档摘要"""
    if doc_id not in documents_store:
        raise HTTPException(status_code=404, detail="文档不存在")
    
    doc = documents_store[doc_id]
    
    # 如果已有摘要且不需要重新生成，直接返回
    if doc.get("summary") and not regenerate:
        print(f"[Summary] 使用缓存的摘要")
        return {"summary": doc["summary"]}
    
    # 重新生成摘要（使用 RAG 检索）
    print(f"[Summary] 重新生成摘要")
    summary = await llm_service.summarize(
        content=doc["content"], 
        doc_type=doc.get("type", "text"),
        doc_id=doc_id,
        use_rag=True  # 启用 RAG 摘要
    )
    
    # 更新存储的摘要
    documents_store[doc_id]["summary"] = summary
    
    return {"summary": summary}


@app.post("/outline/{doc_id}")
async def generate_outline(doc_id: str):
    """生成文档大纲（使用 RAG 优化）"""
    if doc_id not in documents_store:
        raise HTTPException(status_code=404, detail="文档不存在")
    
    doc = documents_store[doc_id]
    outline = await llm_service.generate_outline(
        content=doc["content"],
        doc_id=doc_id,
        use_rag=True
    )
    
    return {"outline": outline}


@app.post("/flashcards/{doc_id}")
async def generate_flashcards(doc_id: str):
    """生成闪卡（使用 RAG 优化）"""
    if doc_id not in documents_store:
        raise HTTPException(status_code=404, detail="文档不存在")
    
    doc = documents_store[doc_id]
    flashcards = await llm_service.generate_flashcards(
        content=doc["content"],
        doc_id=doc_id,
        use_rag=True
    )
    return {"flashcards": flashcards}

@app.post("/quiz/{doc_id}")
async def generate_quiz(doc_id: str):
    """生成测验题目（使用 RAG 优化）"""
    if doc_id not in documents_store:
        raise HTTPException(status_code=404, detail="文档不存在")
    
    doc = documents_store[doc_id]
    quiz = await llm_service.generate_quiz(
        content=doc["content"],
        doc_id=doc_id,
        use_rag=True
    )
    return {"quiz": quiz}

@app.post("/mindmap/{doc_id}")
async def generate_mindmap(doc_id: str):
    """生成思维导图（使用 RAG 优化，服务端渲染 HTML 返回图片）"""
    if doc_id not in documents_store:
        raise HTTPException(status_code=404, detail="文档不存在")
    
    doc = documents_store[doc_id]
    
    # 1. 生成 Markdown 格式的思维导图
    mindmap_markdown = await llm_service.generate_mindmap(
        content=doc["content"],
        doc_id=doc_id,
        use_rag=True
    )
    
    # 2. 使用 Markmap 渲染（专业的思维导图库）
    from app.services.mindmap_markmap_renderer import MindmapMarkmapRenderer
    renderer = MindmapMarkmapRenderer(output_dir="static/mindmaps")
    
    try:
        full_path, relative_path = await renderer.render_to_image(
            markdown=mindmap_markdown,
            width=1920,
            height=1080
        )
        
        return {
            "mindmap": mindmap_markdown,
            "image_url": f"/static/{relative_path}",
            "format": "png"
        }
    except Exception as e:
        import traceback
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=f"思维导图渲染失败: {str(e)}")


class InfographicRequest(BaseModel):
    style: str = "商务简约"
    orientation: str = "landscape"  # landscape, portrait, square
    detailLevel: str = "standard"   # concise, standard, detailed

@app.post("/infographic/{doc_id}")
async def generate_infographic(doc_id: str, request: InfographicRequest = None):
    """生成信息图（使用 RAG 优化）"""
    if doc_id not in documents_store:
        raise HTTPException(status_code=404, detail="文档不存在")
    
    style = request.style if request else "商务简约"
    orientation = request.orientation if request else "landscape"
    detail_level = request.detailLevel if request else "standard"
    
    doc = documents_store[doc_id]
    infographic = await llm_service.generate_infographic(
        content=doc["content"],
        style=style,
        orientation=orientation,
        detail_level=detail_level,
        doc_id=doc_id,
        use_rag=True
    )
    
    return {"infographic": infographic}


# 图片代理接口（解决CORS问题）
@app.get("/proxy-image")
async def proxy_image(url: str):
    """代理外部图片，解决CORS问题"""
    import httpx
    try:
        async with httpx.AsyncClient(timeout=30.0) as client:
            response = await client.get(url)
            if response.status_code == 200:
                content_type = response.headers.get("content-type", "image/png")
                return Response(content=response.content, media_type=content_type)
            else:
                raise HTTPException(status_code=response.status_code, detail="图片获取失败")
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


# 音频概述请求模型
class AudioOverviewRequest(BaseModel):
    format_type: str = "deep_dive"  # deep_dive/brief/critique/debate
    length: str = "default"  # short/default
    focus: str = ""  # 用户指定的关注点


# 报告请求模型
class ReportRequest(BaseModel):
    format_type: str = "blog"  # custom/study_guide/blog
    custom_description: str = ""  # 自己创作时的描述


# 幻灯片请求模型
class SlidesRequest(BaseModel):
    style: str = "business"  # business/tech/academic/creative/minimal
    slide_count: int = 10


# 音频存储
audio_store = {}  # {audio_id: {"file_path": str, "doc_id": str, "title": str}}

# 工作室任务存储
studio_tasks = []  # [{id, type, title, status, audioId?, ...}]


@app.post("/audio-overview/{doc_id}")
async def generate_audio_overview(doc_id: str, request: AudioOverviewRequest = None):
    """生成音频概述"""
    if doc_id not in documents_store:
        raise HTTPException(status_code=404, detail="文档不存在")
    
    doc = documents_store[doc_id]
    format_type = request.format_type if request else "deep_dive"
    length = request.length if request else "default"
    focus = request.focus if request else ""
    
    try:
        # 1. 生成对话脚本
        print(f"[音频] 生成脚本: format={format_type}, length={length}")
        dialogues = await llm_service.generate_audio_script(
            content=doc["content"],
            format_type=format_type,
            length=length,
            focus=focus,
            doc_id=doc_id,
            use_rag=True
        )
        print(f"[音频] 脚本生成完成，共 {len(dialogues)} 段对话")
        
        # 2. 生成音频
        print("[音频] 开始语音合成...")
        audio_file = await tts_service.generate_dialogue(dialogues)
        print(f"[音频] 语音合成完成: {audio_file}")
        
        # 3. 保存音频信息（包含对话脚本）
        import uuid
        audio_id = str(uuid.uuid4())[:8]
        audio_store[audio_id] = {
            "file_path": audio_file,
            "doc_id": doc_id,
            "title": doc["name"],
            "format_type": format_type,
            "dialogues": dialogues  # 保存对话脚本
        }
        
        return {
            "audio_id": audio_id,
            "title": f"{doc['name']} 音频概述",
            "dialogues": dialogues  # 返回脚本用于显示
        }
        
    except Exception as e:
        print(f"[音频] 生成失败: {e}")
        import traceback
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/audio-overview-stream/{doc_id}")
async def generate_audio_overview_stream(doc_id: str, request: AudioOverviewRequest = None):
    """流式生成音频概述（边生成边返回）"""
    if doc_id not in documents_store:
        raise HTTPException(status_code=404, detail="文档不存在")
    
    doc = documents_store[doc_id]
    format_type = request.format_type if request else "deep_dive"
    length = request.length if request else "default"
    focus = request.focus if request else ""
    
    async def audio_stream_generator():
        """流式生成音频数据"""
        try:
            # 1. 生成对话脚本
            print(f"[音频流] 生成脚本: format={format_type}, length={length}")
            dialogues = await llm_service.generate_audio_script(
                content=doc["content"],
                format_type=format_type,
                length=length,
                focus=focus,
                doc_id=doc_id,
                use_rag=True
            )
            print(f"[音频流] 脚本生成完成，共 {len(dialogues)} 段对话")
            
            # 2. 流式生成每段对话的音频
            for i, dialogue in enumerate(dialogues):
                speaker = dialogue.get("speaker", "A")
                text = dialogue.get("text", "")
                
                if not text:
                    continue
                
                voice = "female" if speaker == "A" else "male"
                print(f"[音频流] 生成第 {i+1}/{len(dialogues)} 段: {text[:30]}...")
                
                # 使用流式 TTS
                async for audio_chunk in tts_service.generate_audio_stream(text, voice):
                    # 返回音频数据块
                    yield audio_chunk
                    
            print(f"[音频流] 所有音频片段生成完成")
            
        except Exception as e:
            print(f"[音频流] 生成失败: {e}")
            import traceback
            traceback.print_exc()
            raise
    
    return StreamingResponse(
        audio_stream_generator(),
        media_type="audio/pcm",
        headers={
            "Content-Disposition": f"attachment; filename=podcast_{doc_id}.pcm",
            "X-Content-Type-Options": "nosniff"
        }
    )

@app.get("/audio/{audio_id}")
async def get_audio(audio_id: str):
    """获取音频文件"""
    if audio_id not in audio_store:
        raise HTTPException(status_code=404, detail="音频不存在")
    
    audio_info = audio_store[audio_id]
    file_path = audio_info["file_path"]
    
    if not os.path.exists(file_path):
        raise HTTPException(status_code=404, detail="音频文件不存在")
    
    # 根据文件扩展名确定 media type
    if file_path.endswith('.wav'):
        media_type = "audio/wav"
        filename = f"{audio_info['title']}.wav"
    else:
        media_type = "audio/mpeg"
        filename = f"{audio_info['title']}.mp3"
    
    # 使用 FileResponse，添加 CORS 头
    from urllib.parse import quote
    
    # 对文件名进行 URL 编码以支持中文
    encoded_filename = quote(filename)
    
    response = FileResponse(
        file_path, 
        media_type=media_type,
        filename=filename
    )
    response.headers["Access-Control-Allow-Origin"] = "*"
    return response


@app.get("/audio/{audio_id}/dialogues")
async def get_audio_dialogues(audio_id: str):
    """获取音频对话脚本"""
    if audio_id not in audio_store:
        raise HTTPException(status_code=404, detail="音频不存在")
    
    audio_info = audio_store[audio_id]
    return {
        "dialogues": audio_info.get("dialogues", [])
    }


@app.post("/report/{doc_id}")
async def generate_report(doc_id: str, request: ReportRequest = None):
    """生成报告"""
    if doc_id not in documents_store:
        raise HTTPException(status_code=404, detail="文档不存在")
    
    doc = documents_store[doc_id]
    format_type = request.format_type if request else "blog"
    custom_description = request.custom_description if request else ""
    
    try:
        print(f"[报告] 生成报告: format={format_type}（使用 RAG 优化）")
        content = await llm_service.generate_report(
            content=doc["content"],
            format_type=format_type,
            custom_description=custom_description,
            doc_id=doc_id,
            use_rag=True
        )
        print(f"[报告] 报告生成完成")
        
        return {
            "content": content,
            "title": f"{doc['name']} - 报告",
            "format_type": format_type
        }
        
    except Exception as e:
        print(f"[报告] 生成失败: {e}")
        import traceback
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=str(e))


# ========== PPT 项目 API（LandPPT 风格）==========

from app.services.ppt_service import get_ppt_service

# 获取 PPT 服务实例
def get_ppt_svc():
    return get_ppt_service(llm_service)


class PPTProjectRequest(BaseModel):
    scenario: str = "general"  # general/business/technology/education/academic


class PPTOutlineRequest(BaseModel):
    page_count: int = 10
    target_audience: str = "普通大众"
    ppt_style: str = "general"


class PPTOutlineUpdateRequest(BaseModel):
    outline: dict


@app.post("/ppt/projects/{doc_id}")
async def create_ppt_project(doc_id: str, request: PPTProjectRequest = Body(default=None)):
    """创建 PPT 项目"""
    if doc_id not in documents_store:
        raise HTTPException(status_code=404, detail="文档不存在")
    
    doc = documents_store[doc_id]
    scenario = request.scenario if request else "general"
    
    try:
        ppt_svc = get_ppt_svc()
        project = await ppt_svc.create_project(doc["name"], doc["content"], scenario, doc_id)
        
        return {
            "project_id": project.project_id,
            "topic": project.topic,
            "scenario": project.scenario,
            "status": project.status
        }
    except Exception as e:
        print(f"[PPT] 创建项目失败: {e}")
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/ppt/projects/{project_id}/outline")
async def generate_ppt_outline(project_id: str, request: PPTOutlineRequest = Body(default=None)):
    """生成 PPT 大纲"""
    try:
        ppt_svc = get_ppt_svc()
        page_count = request.page_count if request else 10
        target_audience = request.target_audience if request else "普通大众"
        ppt_style = request.ppt_style if request else "general"
        
        outline = await ppt_svc.generate_outline(
            project_id, 
            page_count=page_count,
            target_audience=target_audience,
            ppt_style=ppt_style
        )
        
        return {
            "status": "success",
            "outline": outline
        }
    except ValueError as e:
        raise HTTPException(status_code=404, detail=str(e))
    except Exception as e:
        print(f"[PPT] 大纲生成失败: {e}")
        import traceback
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=str(e))


@app.put("/ppt/projects/{project_id}/outline")
async def update_ppt_outline(project_id: str, request: PPTOutlineUpdateRequest):
    """更新 PPT 大纲（用户编辑后保存）"""
    try:
        ppt_svc = get_ppt_svc()
        outline = await ppt_svc.update_outline(project_id, request.outline)
        
        return {
            "status": "success",
            "outline": outline
        }
    except ValueError as e:
        raise HTTPException(status_code=404, detail=str(e))
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/ppt/projects/{project_id}/generate-content")
async def generate_ppt_content(project_id: str):
    """生成所有幻灯片的详细内容"""
    try:
        ppt_svc = get_ppt_svc()
        slides = await ppt_svc.generate_all_slides(project_id)
        
        return {
            "status": "success",
            "slides": slides
        }
    except ValueError as e:
        raise HTTPException(status_code=404, detail=str(e))
    except Exception as e:
        print(f"[PPT] 内容生成失败: {e}")
        raise HTTPException(status_code=500, detail=str(e))


@app.get("/ppt/projects/{project_id}/html")
async def get_ppt_html(project_id: str):
    """获取 PPT 的 HTML 预览"""
    try:
        ppt_svc = get_ppt_svc()
        html = ppt_svc.generate_slides_html(project_id)
        
        return Response(content=html, media_type="text/html")
    except ValueError as e:
        raise HTTPException(status_code=404, detail=str(e))
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@app.get("/ppt/projects/{project_id}")
async def get_ppt_project(project_id: str):
    """获取 PPT 项目详情"""
    ppt_svc = get_ppt_svc()
    project = ppt_svc.get_project(project_id)
    
    if not project:
        raise HTTPException(status_code=404, detail="项目不存在")
    
    return {
        "project_id": project.project_id,
        "topic": project.topic,
        "scenario": project.scenario,
        "status": project.status,
        "outline": project.outline,
        "created_at": project.created_at.isoformat(),
        "updated_at": project.updated_at.isoformat()
    }


# ========== Presentation AI 风格 API ==========

from app.services.presentation_service import get_presentation_service, THEMES

def get_pres_svc():
    return get_presentation_service(llm_service)


class PresentationCreateRequest(BaseModel):
    theme: str = "daktilo"
    image_source: str = "pexels"  # pexels 或 ai
    ai_style: str = ""  # AI 生成时的自定义风格


class PresentationOutlineRequest(BaseModel):
    num_slides: int = 8


# 图片来源选项
IMAGE_SOURCES = [
    {"id": "pexels", "name": "Pexels", "description": "免费高质量图片（推荐）"},
    {"id": "ai", "name": "AI 生成", "description": "使用豆包 AI 生成图片"}
]


@app.get("/presentation/themes")
async def get_presentation_themes():
    """获取可用主题列表"""
    return {
        "themes": [
            {"id": k, "name": v["name"], "description": v["description"]}
            for k, v in THEMES.items()
        ]
    }


@app.get("/presentation/image-sources")
async def get_image_sources():
    """获取可用图片来源"""
    return {"sources": IMAGE_SOURCES}


@app.post("/presentation/projects/{doc_id}")
async def create_presentation_project(doc_id: str, request: PresentationCreateRequest = Body(default=None)):
    """创建演示文稿项目"""
    if doc_id not in documents_store:
        raise HTTPException(status_code=404, detail="文档不存在")
    
    doc = documents_store[doc_id]
    theme = request.theme if request else "daktilo"
    image_source = request.image_source if request else "pexels"
    ai_style = request.ai_style if request else ""
    
    print(f"[DEBUG] 创建项目 - theme: {theme}, image_source: {image_source}, ai_style: {ai_style}")
    
    try:
        pres_svc = get_pres_svc()
        project = await pres_svc.create_project(doc["name"], doc["content"], theme, image_source, ai_style, doc_id)
        
        return {
            "project_id": project.project_id,
            "topic": project.topic,
            "theme": project.theme,
            "image_source": project.image_source,
            "status": project.status
        }
    except Exception as e:
        print(f"[Presentation] 创建项目失败: {e}")
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/presentation/projects/{project_id}/outline")
async def generate_presentation_outline(project_id: str, request: PresentationOutlineRequest = Body(default=None)):
    """生成演示文稿大纲"""
    try:
        pres_svc = get_pres_svc()
        num_slides = request.num_slides if request else 8
        
        result = await pres_svc.generate_outline(project_id, num_slides)
        
        return {
            "status": "success",
            "title": result["title"],
            "outline": result["outline"]
        }
    except ValueError as e:
        raise HTTPException(status_code=404, detail=str(e))
    except Exception as e:
        print(f"[Presentation] 大纲生成失败: {e}")
        import traceback
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/presentation/projects/{project_id}/slides")
async def generate_presentation_slides(project_id: str):
    """生成演示文稿幻灯片"""
    try:
        pres_svc = get_pres_svc()
        slides = await pres_svc.generate_slides(project_id)
        
        return {
            "status": "success",
            "slides": slides
        }
    except ValueError as e:
        raise HTTPException(status_code=404, detail=str(e))
    except Exception as e:
        print(f"[Presentation] 幻灯片生成失败: {e}")
        import traceback
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=str(e))


@app.get("/presentation/projects/{project_id}/html")
async def get_presentation_html(project_id: str):
    """获取演示文稿 HTML 预览"""
    try:
        pres_svc = get_pres_svc()
        html = pres_svc.generate_html(project_id)
        
        return Response(content=html, media_type="text/html")
    except ValueError as e:
        raise HTTPException(status_code=404, detail=str(e))
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@app.get("/presentation/projects/{project_id}")
async def get_presentation_project(project_id: str):
    """获取演示文稿项目详情"""
    pres_svc = get_pres_svc()
    project = pres_svc.get_project(project_id)
    
    if not project:
        raise HTTPException(status_code=404, detail="项目不存在")
    
    return {
        "project_id": project.project_id,
        "topic": project.topic,
        "title": project.title,
        "theme": project.theme,
        "status": project.status,
        "outline": project.outline,
        "slides": project.slides
    }


@app.get("/presentation/projects/{project_id}/pptx")
async def download_presentation_pptx(project_id: str):
    """下载演示文稿 PPTX 文件"""
    from pptx import Presentation as PptxPresentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    import io
    
    pres_svc = get_pres_svc()
    project = pres_svc.get_project(project_id)
    
    if not project:
        raise HTTPException(status_code=404, detail="项目不存在")
    
    if not project.slides:
        raise HTTPException(status_code=400, detail="幻灯片尚未生成")
    
    try:
        # 创建 PPTX
        prs = PptxPresentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)
        
        # 主题颜色
        primary_color = RGBColor(59, 130, 246)  # 蓝色
        heading_color = RGBColor(30, 58, 95)    # 深蓝
        text_color = RGBColor(55, 65, 81)       # 灰色
        
        for idx, slide_data in enumerate(project.slides):
            slide_layout = prs.slide_layouts[6]  # 空白布局
            slide = prs.slides.add_slide(slide_layout)
            
            # 获取标题
            slide_title = slide_data.get('title', '') if isinstance(slide_data, dict) else ''
            
            # 添加标题
            title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11.5), Inches(1))
            title_frame = title_box.text_frame
            title_para = title_frame.paragraphs[0]
            title_para.text = slide_title
            title_para.font.size = Pt(40)
            title_para.font.bold = True
            title_para.font.color.rgb = heading_color
            
            # 添加内容 - 从 elements 中提取
            elements = slide_data.get('elements', []) if isinstance(slide_data, dict) else []
            content_top = 1.6
            
            for elem in elements:
                if isinstance(elem, dict):
                    # 添加元素标题
                    elem_title = elem.get('title', '')
                    if elem_title:
                        box = slide.shapes.add_textbox(Inches(0.8), Inches(content_top), Inches(11.5), Inches(0.5))
                        frame = box.text_frame
                        p = frame.paragraphs[0]
                        p.text = elem_title
                        p.font.size = Pt(22)
                        p.font.bold = True
                        p.font.color.rgb = primary_color
                        content_top += 0.5
                    
                    # 添加元素内容
                    elem_content = elem.get('content', '')
                    if elem_content:
                        box = slide.shapes.add_textbox(Inches(0.8), Inches(content_top), Inches(11.5), Inches(0.8))
                        frame = box.text_frame
                        frame.word_wrap = True
                        p = frame.paragraphs[0]
                        p.text = elem_content
                        p.font.size = Pt(16)
                        p.font.color.rgb = text_color
                        content_top += 0.6
                    
                    # 添加列表项
                    items = elem.get('items', [])
                    for item in items:
                        box = slide.shapes.add_textbox(Inches(1.0), Inches(content_top), Inches(11), Inches(0.4))
                        frame = box.text_frame
                        frame.word_wrap = True
                        p = frame.paragraphs[0]
                        p.text = f"• {item}"
                        p.font.size = Pt(14)
                        p.font.color.rgb = text_color
                        content_top += 0.35
                    
                    content_top += 0.2
            
            # 添加页码
            page_box = slide.shapes.add_textbox(Inches(12), Inches(7), Inches(1), Inches(0.3))
            page_frame = page_box.text_frame
            page_para = page_frame.paragraphs[0]
            page_para.text = f"{idx + 1} / {len(project.slides)}"
            page_para.font.size = Pt(11)
            page_para.font.color.rgb = RGBColor(150, 150, 150)
        
        # 保存到内存
        pptx_io = io.BytesIO()
        prs.save(pptx_io)
        pptx_io.seek(0)
        
        from fastapi.responses import StreamingResponse
        
        title = project.title or '演示文稿'
        import urllib.parse
        encoded_title = urllib.parse.quote(title)
        
        return StreamingResponse(
            pptx_io,
            media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
            headers={"Content-Disposition": f"attachment; filename*=UTF-8''{encoded_title}.pptx"}
        )
    except Exception as e:
        print(f"生成 PPTX 失败: {e}")
        import traceback
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=f"生成 PPTX 失败: {str(e)}")


@app.get("/presentation/projects/{project_id}/pdf")
async def download_presentation_pdf(project_id: str):
    """下载 AI 生成的幻灯片为 PDF（图片合集）"""
    from reportlab.lib.pagesizes import landscape, A4
    from reportlab.pdfgen import canvas
    from reportlab.lib.utils import ImageReader
    import io
    import httpx
    
    pres_svc = get_pres_svc()
    project = pres_svc.get_project(project_id)
    
    if not project:
        raise HTTPException(status_code=404, detail="项目不存在")
    
    if not project.slides:
        raise HTTPException(status_code=400, detail="幻灯片尚未生成")
    
    try:
        # 创建 PDF
        pdf_io = io.BytesIO()
        page_width, page_height = landscape(A4)  # 横向 A4
        c = canvas.Canvas(pdf_io, pagesize=(page_width, page_height))
        
        async with httpx.AsyncClient(timeout=60.0) as client:
            for idx, slide_data in enumerate(project.slides):
                # 获取 AI 生成的幻灯片图片 URL
                image_url = slide_data.get('slide_image_url') if isinstance(slide_data, dict) else None
                
                if image_url:
                    try:
                        # 下载图片
                        response = await client.get(image_url)
                        if response.status_code == 200:
                            img_data = io.BytesIO(response.content)
                            img = ImageReader(img_data)
                            
                            # 计算图片尺寸以适应页面（保持比例）
                            img_width, img_height = img.getSize()
                            ratio = min(page_width / img_width, page_height / img_height)
                            new_width = img_width * ratio
                            new_height = img_height * ratio
                            
                            # 居中绘制
                            x = (page_width - new_width) / 2
                            y = (page_height - new_height) / 2
                            
                            c.drawImage(img, x, y, width=new_width, height=new_height)
                    except Exception as e:
                        print(f"下载幻灯片 {idx + 1} 图片失败: {e}")
                        # 绘制占位文字
                        c.setFont("Helvetica", 24)
                        c.drawCentredString(page_width / 2, page_height / 2, f"Slide {idx + 1}")
                else:
                    # 没有图片，绘制标题
                    slide_title = slide_data.get('title', f'Slide {idx + 1}') if isinstance(slide_data, dict) else f'Slide {idx + 1}'
                    c.setFont("Helvetica", 24)
                    c.drawCentredString(page_width / 2, page_height / 2, slide_title)
                
                # 添加新页面（除了最后一页）
                if idx < len(project.slides) - 1:
                    c.showPage()
        
        c.save()
        pdf_io.seek(0)
        
        from fastapi.responses import StreamingResponse
        
        title = project.title or '演示文稿'
        import urllib.parse
        encoded_title = urllib.parse.quote(title)
        
        return StreamingResponse(
            pdf_io,
            media_type="application/pdf",
            headers={"Content-Disposition": f"attachment; filename*=UTF-8''{encoded_title}.pdf"}
        )
    except Exception as e:
        print(f"生成 PDF 失败: {e}")
        import traceback
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=f"生成 PDF 失败: {str(e)}")


# 网页导入接口
class URLImportRequest(BaseModel):
    url: str

@app.post("/import-url", response_model=DocumentInfo)
async def import_url(request: URLImportRequest):
    """导入网页内容（增强版：支持 JS 渲染、智能提取、反爬虫）"""
    
    url = request.url
    
    try:
        # 使用增强版爬虫服务
        result = await web_scraper_service.scrape(url, use_js=True)
        
        # 检查是否成功
        if not result["success"]:
            error_msg = result.get("error", "未知错误")
            raise HTTPException(
                status_code=400,
                detail=f"网页爬取失败: {error_msg}。可能原因：1) 网页需要登录 2) 网页被反爬虫保护 3) 网络连接问题"
            )
        
        title_text = result["title"]
        text = result["content"]
        method = result["method"]
        
        print(f"[Web] 爬取成功 (方法: {method})，标题: {title_text}, 内容长度: {len(text)}")
        
        # 验证内容是否为空
        if not text or len(text.strip()) < 50:
            raise HTTPException(
                status_code=400,
                detail=f"网页内容为空或过短（仅 {len(text)} 字符），可能是：1) 网页需要登录 2) 内容由JavaScript动态加载 3) 网页被反爬虫保护"
            )
        
        # 生成文档 ID（使用UUID避免冲突）
        doc_id = f"doc_{uuid.uuid4().hex[:12]}"
        
        # 存储文档
        documents_store[doc_id] = {
            "id": doc_id,
            "name": title_text,
            "type": "web",
            "content": text,
            "url": url
        }
        
        return DocumentInfo(
            id=doc_id,
            name=title_text,
            type="web",
            content_preview=text[:200] + "..." if len(text) > 200 else text,
            url=url
        )
        
    except Exception as e:
        print(f"网页导入失败: {e}")
        import traceback
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=f"网页导入失败: {str(e)}")


# 粘贴文本接口
class PasteTextRequest(BaseModel):
    text: str
    title: str = None

@app.post("/paste-text", response_model=DocumentInfo)
async def paste_text(request: PasteTextRequest):
    """粘贴文本内容"""
    
    text = request.text.strip()
    
    if not text:
        raise HTTPException(status_code=400, detail="文本内容不能为空")
    
    # 生成文档 ID
    doc_id = f"doc_{uuid.uuid4().hex[:12]}"
    
    # 生成标题（如果没有提供）
    if request.title:
        title = request.title.strip()
    else:
        # 使用文本前 30 个字符作为标题
        title = text[:30] + "..." if len(text) > 30 else text
    
    print(f"[Paste] 粘贴文本，标题: {title}, 长度: {len(text)}")
    
    # 存储文档
    documents_store[doc_id] = {
        "id": doc_id,
        "name": title,
        "type": "text",
        "content": text
    }
    
    # 生成摘要
    try:
        summary = await llm_service.generate_summary(text)
        documents_store[doc_id]["summary"] = summary
    except Exception as e:
        print(f"[Paste] 生成摘要失败: {e}")
        documents_store[doc_id]["summary"] = text[:200] + "..." if len(text) > 200 else text
    
    # 添加到向量库
    try:
        await rag_service.add_document(doc_id, title, text)
        print(f"[Paste] 已添加到向量库")
    except Exception as e:
        print(f"[Paste] 添加到向量库失败: {e}")
    
    return DocumentInfo(
        id=doc_id,
        name=title,
        type="text",
        content_preview=text[:200] + "..." if len(text) > 200 else text
    )


# Word 预览接口
@app.get("/documents/{doc_id}/word-preview")
async def preview_word(doc_id: str):
    """预览 Word 文件内容（返回 HTML）"""
    if doc_id not in documents_store:
        raise HTTPException(status_code=404, detail="文档不存在")
    
    doc = documents_store[doc_id]
    if doc["type"] != "word":
        raise HTTPException(status_code=400, detail="不是 Word 文件")
    
    # 转换 Word 为 HTML
    try:
        import mammoth
        import base64
        from io import BytesIO
        
        # 从 base64 还原为二进制
        content = base64.b64decode(doc["raw_content"])
        
        # 使用 mammoth 转换为 HTML
        result = mammoth.convert_to_html(BytesIO(content))
        html_content = result.value
        
        return {"html": html_content, "name": doc.get("name", "")}
    except Exception as e:
        print(f"Word 预览失败: {e}")
        import traceback
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=f"Word 预览失败: {str(e)}")

# Excel 预览接口
@app.get("/documents/{doc_id}/excel-preview")
async def preview_excel(doc_id: str):
    """预览 Excel 文件内容（返回结构化数据）"""
    if doc_id not in documents_store:
        raise HTTPException(status_code=404, detail="文档不存在")
    
    doc = documents_store[doc_id]
    if doc["type"] != "excel":
        raise HTTPException(status_code=400, detail="不是 Excel 文件")
    
    # 解析 Excel 为 JSON
    try:
        import pandas as pd
        from io import BytesIO
        import base64
        
        print(f"[DEBUG] 开始解析 Excel, doc_id: {doc_id}")
        
        # 从 base64 还原为二进制
        content = base64.b64decode(doc["raw_content"])
        print(f"[DEBUG] 文件大小: {len(content)} bytes")
        
        # 使用 pandas 读取 Excel
        excel_file = pd.ExcelFile(BytesIO(content))
        print(f"[DEBUG] 工作表列表: {excel_file.sheet_names}")
        
        # 解析所有工作表
        sheets_data = {}
        for sheet_name in excel_file.sheet_names:
            df = pd.read_excel(excel_file, sheet_name=sheet_name)
            print(f"[DEBUG] 工作表 '{sheet_name}' 形状: {df.shape}")
            print(f"[DEBUG] 工作表 '{sheet_name}' 列名: {df.columns.tolist()}")
            print(f"[DEBUG] 工作表 '{sheet_name}' 前3行:\n{df.head(3)}")
            # 转换为字典列表
            sheets_data[sheet_name] = df.fillna('').to_dict('records')
            print(f"[DEBUG] 工作表 '{sheet_name}' 转换后数据量: {len(sheets_data[sheet_name])}")
        
        return sheets_data
    except Exception as e:
        print(f"[ERROR] Excel 预览失败: {e}")
        import traceback
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=f"Excel 预览失败: {str(e)}")

# 保留旧的简单接口用于兼容
@app.post("/slides/{doc_id}")
async def generate_slides(doc_id: str, request: SlidesRequest = Body(default=None)):
    """生成幻灯片（简单接口）"""
    if doc_id not in documents_store:
        raise HTTPException(status_code=404, detail="文档不存在")
    
    doc = documents_store[doc_id]
    style = request.style if request else "business"
    slide_count = request.slide_count if request else 10
    
    try:
        print(f"[幻灯片] 生成幻灯片: style={style}, count={slide_count}")
        
        # 调用 LLM 生成幻灯片内容
        slides = await llm_service.generate_slides(doc["content"], style, slide_count)
        
        print(f"[幻灯片] 生成完成，共 {len(slides)} 页")
        
        return {
            "slides": slides,
            "title": doc["name"],
            "style": style
        }
        
    except Exception as e:
        print(f"[幻灯片] 生成失败: {e}")
        import traceback
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=str(e))


class SlidesDownloadRequest(BaseModel):
    slides: list


@app.post("/slides/{doc_id}/download")
async def download_slides_pptx(doc_id: str, request: SlidesDownloadRequest, style: str = "business"):
    """生成并下载 PPT 文件（预留接口，等待开源模型对接）"""
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RgbColor
    import io
    
    slides = request.slides
    if not slides:
        raise HTTPException(status_code=400, detail="幻灯片内容为空")
    
    # 获取文档标题
    title = documents_store.get(doc_id, {}).get("name", "幻灯片")
    
    # 创建 PPT
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # 风格配置
    style_config = {
        "business": {"bg": RgbColor(30, 58, 95), "title": RgbColor(255, 255, 255), "text": RgbColor(200, 200, 200)},
        "tech": {"bg": RgbColor(15, 23, 42), "title": RgbColor(34, 211, 238), "text": RgbColor(200, 200, 200)},
        "academic": {"bg": RgbColor(250, 250, 245), "title": RgbColor(30, 30, 30), "text": RgbColor(60, 60, 60)},
        "creative": {"bg": RgbColor(251, 207, 232), "title": RgbColor(190, 24, 93), "text": RgbColor(80, 80, 80)},
        "minimal": {"bg": RgbColor(255, 255, 255), "title": RgbColor(30, 30, 30), "text": RgbColor(80, 80, 80)}
    }
    config = style_config.get(style, style_config["business"])
    
    for slide_data in slides:
        # 使用空白布局
        blank_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(blank_layout)
        
        # 设置背景色
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = config["bg"]
        
        # 添加标题
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(11.7), Inches(1.5))
        title_frame = title_box.text_frame
        title_frame.word_wrap = True
        p = title_frame.paragraphs[0]
        p.text = slide_data.get("title", "")
        p.font.size = Pt(44)
        p.font.bold = True
        p.font.color.rgb = config["title"]
        
        # 添加要点
        points = slide_data.get("points", [])
        if points:
            content_box = slide.shapes.add_textbox(Inches(0.8), Inches(2.5), Inches(11.7), Inches(4.5))
            content_frame = content_box.text_frame
            content_frame.word_wrap = True
            
            for i, point in enumerate(points):
                if i == 0:
                    p = content_frame.paragraphs[0]
                else:
                    p = content_frame.add_paragraph()
                p.text = f"• {point}"
                p.font.size = Pt(28)
                p.font.color.rgb = config["text"]
                p.space_after = Pt(16)
    
    # 保存到内存
    pptx_io = io.BytesIO()
    prs.save(pptx_io)
    pptx_io.seek(0)
    
    from fastapi.responses import StreamingResponse
    
    return StreamingResponse(
        pptx_io,
        media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        headers={"Content-Disposition": f"attachment; filename={title}.pptx"}
    )


# 视频请求模型
class VideoRequest(BaseModel):
    duration: int = 60  # 视频时长（秒）
    style: str = "professional"  # professional/casual/educational


# 视频存储
video_store = {}  # {video_id: {"file_path": str, "doc_id": str, "title": str}}




@app.get("/video/{video_id}")
async def get_video(video_id: str):
    """获取视频文件"""
    from fastapi.responses import FileResponse
    
    if video_id not in video_store:
        raise HTTPException(status_code=404, detail="视频不存在")
    
    video_info = video_store[video_id]
    file_path = video_info["file_path"]
    
    if not os.path.exists(file_path):
        raise HTTPException(status_code=404, detail="视频文件不存在")
    
    return FileResponse(
        file_path,
        media_type="video/mp4",
        filename=f"{video_info['title']}.mp4"
    )


# ========== 工作室任务 API ==========

class StudioTask(BaseModel):
    id: str
    type: str  # audio, mindmap, faq, timeline, etc.
    title: str
    status: str  # pending, completed, error
    audioId: Optional[str] = None
    data: Optional[dict] = None


@app.get("/studio/tasks")
async def get_studio_tasks():
    """获取所有工作室任务"""
    return {"tasks": studio_tasks}


@app.post("/studio/tasks")
async def add_studio_task(task: StudioTask):
    """添加工作室任务"""
    # 检查是否已存在
    for i, t in enumerate(studio_tasks):
        if t["id"] == task.id:
            studio_tasks[i] = task.dict()
            return {"success": True, "task": task}
    
    studio_tasks.append(task.dict())
    return {"success": True, "task": task}


@app.put("/studio/tasks/{task_id}")
async def update_studio_task(task_id: str, task: StudioTask):
    """更新工作室任务"""
    for i, t in enumerate(studio_tasks):
        if t["id"] == task_id:
            studio_tasks[i] = task.dict()
            return {"success": True, "task": task}
    
    raise HTTPException(status_code=404, detail="任务不存在")


@app.delete("/studio/tasks/{task_id}")
async def delete_studio_task(task_id: str):
    """删除工作室任务"""
    global studio_tasks
    studio_tasks = [t for t in studio_tasks if t["id"] != task_id]
    return {"success": True}


@app.delete("/studio/tasks")
async def clear_studio_tasks():
    """清空所有工作室任务"""
    global studio_tasks
    studio_tasks = []
    return {"success": True}


# ==================== 会话管理 API ====================

class ConversationCreate(BaseModel):
    title: str
    document_ids: List[str] = []


class MessageCreate(BaseModel):
    conversation_id: str
    role: str
    content: str
    sources: Optional[List[dict]] = None
    suggestions: Optional[List[str]] = None
    citation_map: Optional[dict] = None


@app.post("/conversations")
async def create_conversation(conv: ConversationCreate, db: Session = Depends(get_db)):
    """创建新会话"""
    conversation = Conversation(
        id=f"conv_{uuid.uuid4().hex[:12]}",
        title=conv.title
    )
    db.add(conversation)
    db.commit()
    db.refresh(conversation)
    
    return {
        "id": conversation.id,
        "title": conversation.title,
        "created_at": conversation.created_at.isoformat(),
        "updated_at": conversation.updated_at.isoformat()
    }


@app.get("/conversations")
async def get_conversations(db: Session = Depends(get_db)):
    """获取所有会话列表"""
    conversations = db.query(Conversation).order_by(Conversation.updated_at.desc()).all()
    return {
        "conversations": [
            {
                "id": conv.id,
                "title": conv.title,
                "created_at": conv.created_at.isoformat(),
                "updated_at": conv.updated_at.isoformat(),
                "message_count": len(conv.messages)
            }
            for conv in conversations
        ]
    }


@app.get("/conversations/{conversation_id}")
async def get_conversation(conversation_id: str, db: Session = Depends(get_db)):
    """获取指定会话的所有消息"""
    conversation = db.query(Conversation).filter(Conversation.id == conversation_id).first()
    if not conversation:
        raise HTTPException(status_code=404, detail="会话不存在")
    
    messages = db.query(Message).filter(Message.conversation_id == conversation_id).order_by(Message.created_at).all()
    
    return {
        "id": conversation.id,
        "title": conversation.title,
        "created_at": conversation.created_at.isoformat(),
        "updated_at": conversation.updated_at.isoformat(),
        "messages": [
            {
                "id": msg.id,
                "role": msg.role,
                "content": msg.content,
                "sources": msg.sources,
                "suggestions": msg.suggestions,
                "citationMap": msg.citation_map,
                "created_at": msg.created_at.isoformat()
            }
            for msg in messages
        ]
    }


@app.post("/conversations/{conversation_id}/messages")
async def add_message(conversation_id: str, msg: MessageCreate, db: Session = Depends(get_db)):
    """向会话添加消息"""
    conversation = db.query(Conversation).filter(Conversation.id == conversation_id).first()
    if not conversation:
        raise HTTPException(status_code=404, detail="会话不存在")
    
    message = Message(
        id=f"msg_{uuid.uuid4().hex[:12]}",
        conversation_id=conversation_id,
        role=msg.role,
        content=msg.content,
        sources=msg.sources,
        suggestions=msg.suggestions,
        citation_map=msg.citation_map
    )
    db.add(message)
    
    # 更新会话的 updated_at
    conversation.updated_at = datetime.utcnow()
    
    db.commit()
    db.refresh(message)
    
    return {
        "id": message.id,
        "role": message.role,
        "content": message.content,
        "sources": message.sources,
        "suggestions": message.suggestions,
        "citationMap": message.citation_map,
        "created_at": message.created_at.isoformat()
    }


@app.delete("/conversations/{conversation_id}")
async def delete_conversation(conversation_id: str, db: Session = Depends(get_db)):
    """删除会话"""
    conversation = db.query(Conversation).filter(Conversation.id == conversation_id).first()
    if not conversation:
        raise HTTPException(status_code=404, detail="会话不存在")
    
    db.delete(conversation)
    db.commit()
    
    return {"success": True}


if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=8000)
