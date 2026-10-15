import fitz  # PyMuPDF
from docx import Document
from io import BytesIO
import base64


class DocumentParser:
    """文档解析服务"""
    
    def parse_pdf(self, content: bytes) -> dict:
        """解析 PDF 文件，返回文本和需要 OCR 的图片"""
        try:
            doc = fitz.open(stream=content, filetype="pdf")
            text = ""
            images = []  # 需要视觉模型 OCR 的页面
            total_pages = len(doc)
            text_pages = 0
            
            for page_num, page in enumerate(doc, 1):
                # 先尝试提取文字
                page_text = page.get_text()
                
                # 检测是否为乱码
                is_garbled = self._is_garbled_text(page_text)
                
                if page_text.strip() and len(page_text.strip()) > 50 and not is_garbled:
                    # 有足够的文字内容且不是乱码
                    text += f"\n[第 {page_num} 页]\n{page_text}"
                    text_pages += 1
                else:
                    # 文字太少、是乱码、或可能是扫描件，转成图片供视觉模型识别
                    if is_garbled:
                        print(f"[Parser] 第 {page_num} 页检测到乱码，需要 OCR")
                    img_base64 = self._page_to_base64(page)
                    if img_base64:
                        images.append({
                            "page": page_num,
                            "base64": img_base64
                        })
                        print(f"[Parser] 第 {page_num} 页需要 OCR")
            
            doc.close()
            
            # 计算文字页占比
            text_ratio = text_pages / total_pages if total_pages > 0 else 0
            
            return {
                "text": text.strip(),
                "images": images,
                "total_pages": total_pages,
                "text_pages": text_pages,
                "text_ratio": text_ratio,
                "quality": self._assess_quality(text, len(content), text_ratio)
            }
        except Exception as e:
            raise Exception(f"PDF 解析失败: {str(e)}")
    
    def _assess_quality(self, text: str, file_size: int, text_ratio: float) -> str:
        """评估解析质量（多算法综合）"""
        if not text or len(text.strip()) < 100:
            return "poor"
        
        score = 0  # 质量评分（0-100）
        
        # 1. 文字页占比（权重：30%）
        if text_ratio > 0.8:
            score += 30
        elif text_ratio > 0.5:
            score += 20
        elif text_ratio > 0.3:
            score += 10
        else:
            score += 0  # < 30% 很可能是扫描件
        
        # 2. 文字密度（权重：25%）
        file_size_mb = file_size / (1024 * 1024)
        char_count = len(text)
        density = char_count / file_size_mb if file_size_mb > 0 else 0
        
        if density > 2000:
            score += 25
        elif density > 1000:
            score += 20
        elif density > 500:
            score += 10
        else:
            score += 0
        
        # 3. 乱码检测（权重：20%）
        if not self._detect_garbled_text(text):
            score += 20
        
        # 4. 语言连贯性（权重：15%）
        coherence = self._check_language_coherence(text)
        score += int(coherence * 15)
        
        # 5. 重复内容检测（权重：10%）
        repetition = self._detect_repetition(text)
        if repetition < 0.3:  # 重复率低于 30%
            score += 10
        elif repetition < 0.5:
            score += 5
        
        # 综合判断
        if score >= 70:
            return "good"
        elif score >= 40:
            return "medium"
        else:
            return "poor"
    
    def _is_garbled_text(self, text: str) -> bool:
        """检测单页文本是否为乱码（用于实时判断）"""
        return self._detect_garbled_text(text)
    
    def _detect_garbled_text(self, text: str) -> bool:
        """检测是否为乱码"""
        if not text or len(text) < 50:
            return True
        
        # 统计字符类型
        chinese_count = sum(1 for c in text if '\u4e00' <= c <= '\u9fff')
        english_count = sum(1 for c in text if c.isalpha() and ord(c) < 128)
        digit_count = sum(1 for c in text if c.isdigit())
        special_count = sum(1 for c in text if not c.isalnum() and not c.isspace())
        
        total_chars = len(text.replace(' ', '').replace('\n', ''))
        if total_chars == 0:
            return True
        
        # 正常文本应该有合理的字符分布
        normal_ratio = (chinese_count + english_count + digit_count) / total_chars
        
        # 特殊字符太多可能是乱码
        if special_count / total_chars > 0.5:
            return True
        
        # 正常字符太少也是乱码
        if normal_ratio < 0.5:
            return True
        
        return False
    
    def _check_language_coherence(self, text: str) -> float:
        """检测语言连贯性（0-1）"""
        if not text or len(text) < 100:
            return 0.0
        
        # 提取前 1000 字符
        sample = text[:1000]
        
        # 统计常见词汇
        common_chinese_words = ['的', '是', '在', '了', '和', '有', '我', '你', '他', '这', '那', '个', '人', '中', '大', '小']
        common_english_words = ['the', 'is', 'are', 'and', 'or', 'in', 'on', 'at', 'to', 'of', 'for']
        
        sample_lower = sample.lower()
        
        # 统计常见词出现次数
        chinese_word_count = sum(sample.count(word) for word in common_chinese_words)
        english_word_count = sum(sample_lower.count(' ' + word + ' ') for word in common_english_words)
        
        total_words = len(sample.split())
        if total_words == 0:
            return 0.0
        
        # 常见词占比
        coherence_score = (chinese_word_count + english_word_count) / total_words
        
        return min(coherence_score, 1.0)
    
    def _detect_repetition(self, text: str) -> float:
        """检测重复内容比例（0-1）"""
        if not text or len(text) < 100:
            return 0.0
        
        # 按行分割
        lines = [line.strip() for line in text.split('\n') if line.strip() and len(line.strip()) > 5]
        if len(lines) < 10:
            return 0.0
        
        # 统计重复行
        from collections import Counter
        line_counts = Counter(lines)
        
        # 重复超过 3 次的行
        repeated_lines = sum(count - 1 for line, count in line_counts.items() if count > 3)
        repetition_ratio = repeated_lines / len(lines) if len(lines) > 0 else 0.0
        
        return min(repetition_ratio, 1.0)
    
    def _page_to_base64(self, page, max_size: int = 2000) -> str:
        """将 PDF 页面转换为 base64 图片，自动压缩大图"""
        try:
            # 获取页面尺寸
            rect = page.rect
            width, height = rect.width, rect.height
            
            # 计算缩放比例，确保最长边不超过 max_size
            scale = min(max_size / max(width, height), 2.0)
            mat = fitz.Matrix(scale, scale)
            
            pix = page.get_pixmap(matrix=mat)
            img_bytes = pix.tobytes("png")
            
            # 如果图片太大（>5MB），进一步压缩
            if len(img_bytes) > 5 * 1024 * 1024:
                scale = scale * 0.5
                mat = fitz.Matrix(scale, scale)
                pix = page.get_pixmap(matrix=mat)
                img_bytes = pix.tobytes("png")
            
            return base64.b64encode(img_bytes).decode('utf-8')
        except Exception as e:
            print(f"[Parser] 页面转图片失败: {e}")
            return None
    
    def parse_docx(self, content: bytes) -> dict:
        """解析 Word 文档（增强版）"""
        try:
            doc = Document(BytesIO(content))
            paragraphs = []
            tables_text = []
            
            # 1. 提取段落
            for para in doc.paragraphs:
                if para.text.strip():
                    paragraphs.append(para.text)
            
            # 2. 提取表格
            for table in doc.tables:
                table_data = []
                for row in table.rows:
                    row_data = [cell.text.strip() for cell in row.cells]
                    if any(row_data):  # 至少有一个非空单元格
                        table_data.append(" | ".join(row_data))
                if table_data:
                    tables_text.append("\n".join(table_data))
            
            # 3. 合并文本
            text = "\n\n".join(paragraphs)
            if tables_text:
                text += "\n\n[表格内容]\n" + "\n\n".join(tables_text)
            
            # 4. 质量评估
            quality = self._assess_docx_quality(text, len(content))
            
            return {
                "text": text,
                "quality": quality,
                "has_tables": len(tables_text) > 0,
                "paragraph_count": len(paragraphs),
                "table_count": len(tables_text)
            }
        except Exception as e:
            raise Exception(f"Word 文档解析失败: {str(e)}")
    
    def _assess_docx_quality(self, text: str, file_size: int) -> str:
        """评估 Word 解析质量"""
        if not text or len(text.strip()) < 50:
            return "poor"
        
        # 文字密度
        file_size_mb = file_size / (1024 * 1024)
        char_count = len(text)
        density = char_count / file_size_mb if file_size_mb > 0 else 0
        
        # Word 文档通常密度较高
        if density > 1000:
            return "good"
        elif density > 300:
            return "medium"
        else:
            return "poor"  # 可能有大量图片或复杂格式
    
    def parse_excel(self, content: bytes) -> dict:
        """解析 Excel 文档（本地解析）"""
        try:
            import pandas as pd
            from io import BytesIO
            
            # 读取所有 sheet
            excel_file = BytesIO(content)
            sheets = pd.read_excel(excel_file, sheet_name=None, engine='openpyxl')
            
            all_text = []
            total_rows = 0
            
            for sheet_name, df in sheets.items():
                # 跳过空 sheet
                if df.empty:
                    continue
                
                # Sheet 标题
                all_text.append(f"[工作表: {sheet_name}]")
                
                # 转换为文本
                # 方式1：表格格式（保留结构）
                table_text = df.to_string(index=False, na_rep='')
                all_text.append(table_text)
                
                total_rows += len(df)
            
            text = "\n\n".join(all_text)
            
            # 质量评估
            quality = self._assess_excel_quality(text, len(content), total_rows)
            
            return {
                "text": text,
                "quality": quality,
                "sheet_count": len(sheets),
                "total_rows": total_rows
            }
        except Exception as e:
            # 本地解析失败，返回 poor 质量，让上层调用 Qwen-Long
            return {
                "text": "",
                "quality": "poor",
                "error": str(e)
            }
    
    def _assess_excel_quality(self, text: str, file_size: int, total_rows: int) -> str:
        """评估 Excel 解析质量"""
        if not text or len(text.strip()) < 50:
            return "poor"
        
        # Excel 特点：行数多，但文件可能很大（有公式、格式）
        if total_rows > 100:
            return "good"  # 数据量大，本地解析足够
        elif total_rows > 20:
            return "medium"
        else:
            # 行数少但文件大，可能有复杂公式或图表
            file_size_mb = file_size / (1024 * 1024)
            if file_size_mb > 1 and total_rows < 50:
                return "poor"  # 需要 Qwen-Long 理解公式
            else:
                return "medium"
    
    def parse_txt(self, content: bytes) -> dict:
        """解析纯文本文件（智能编码检测）"""
        # 常见编码列表（按优先级）
        encodings = ['utf-8', 'gbk', 'gb2312', 'big5', 'utf-16', 'latin1']
        
        text = None
        detected_encoding = None
        
        # 1. 尝试自动检测编码
        try:
            import chardet
            detected = chardet.detect(content)
            if detected['confidence'] > 0.7:  # 置信度 > 70%
                detected_encoding = detected['encoding']
                try:
                    text = content.decode(detected_encoding)
                    return {
                        "text": text,
                        "encoding": detected_encoding,
                        "confidence": detected['confidence'],
                        "quality": "good"
                    }
                except:
                    pass
        except ImportError:
            pass  # chardet 未安装，使用手动检测
        
        # 2. 手动尝试常见编码
        for encoding in encodings:
            try:
                text = content.decode(encoding)
                # 检测是否有乱码
                if not self._has_mojibake(text):
                    return {
                        "text": text,
                        "encoding": encoding,
                        "confidence": 0.9,
                        "quality": "good"
                    }
            except (UnicodeDecodeError, LookupError):
                continue
        
        # 3. 最后尝试：忽略错误
        text = content.decode('utf-8', errors='ignore')
        return {
            "text": text,
            "encoding": "utf-8 (with errors)",
            "confidence": 0.5,
            "quality": "medium"
        }
    
    def _has_mojibake(self, text: str) -> bool:
        """检测是否有乱码"""
        if not text or len(text) < 10:
            return False
        
        # 检测常见乱码模式
        mojibake_patterns = [
            '�',  # 替换字符
            '\ufffd',  # Unicode 替换字符
        ]
        
        for pattern in mojibake_patterns:
            if pattern in text:
                return True
        
        # 检测异常字符比例
        sample = text[:1000]
        control_chars = sum(1 for c in sample if ord(c) < 32 and c not in '\n\r\t')
        if len(sample) > 0 and control_chars / len(sample) > 0.1:
            return True
        
        return False
    
    def parse_markdown(self, content: bytes) -> dict:
        """解析 Markdown 文件（增强版）"""
        # 1. 先解码
        txt_result = self.parse_txt(content)
        text = txt_result["text"]
        
        # 2. 提取结构信息
        lines = text.split('\n')
        
        # 统计 Markdown 元素
        headers = []
        code_blocks = 0
        links = 0
        images = 0
        tables = 0
        
        in_code_block = False
        for line in lines:
            stripped = line.strip()
            
            # 标题
            if stripped.startswith('#'):
                level = len(stripped) - len(stripped.lstrip('#'))
                title = stripped.lstrip('#').strip()
                if title:
                    headers.append({"level": level, "title": title})
            
            # 代码块
            if stripped.startswith('```'):
                in_code_block = not in_code_block
                if in_code_block:
                    code_blocks += 1
            
            # 链接
            links += line.count('](')
            
            # 图片
            images += line.count('![')
            
            # 表格
            if '|' in line and not in_code_block:
                tables += 1
        
        # 3. 质量评估
        quality = "good"
        if len(headers) == 0 and code_blocks == 0:
            # 可能不是真正的 Markdown，只是普通文本
            quality = "medium"
        
        return {
            "text": text,
            "encoding": txt_result["encoding"],
            "quality": quality,
            "structure": {
                "headers": len(headers),
                "code_blocks": code_blocks,
                "links": links,
                "images": images,
                "tables": tables
            },
            "toc": headers[:20]  # 前 20 个标题作为目录
        }
